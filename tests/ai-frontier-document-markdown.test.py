"""Focused offline conversion tests; run with the sandbox's installed Python dependencies."""

import argparse
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

sys.dont_write_bytecode = True
SCRIPT = Path(__file__).resolve().parents[1] / "skills/document-markdown/convert.py"
SPEC = importlib.util.spec_from_file_location("document_markdown", SCRIPT)
converter = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(converter)
HAS_OPENPYXL = importlib.util.find_spec("openpyxl") is not None
HAS_XLRD = importlib.util.find_spec("xlrd") is not None
HAS_PANDAS = importlib.util.find_spec("pandas") is not None
HAS_CHARDET = importlib.util.find_spec("chardet") is not None
LIBREOFFICE = shutil.which("libreoffice")


def options(**overrides):
    return argparse.Namespace(**{
        "start_row": 1, "start_column": 1, "max_rows": 200, "max_columns": 50,
        "encoding": None, **overrides,
    })


class ConversionTests(unittest.TestCase):
    def setUp(self):
        self.previous = Path.cwd()
        self.temporary = tempfile.TemporaryDirectory(prefix="frontier-document-markdown-test-")
        self.root = Path(self.temporary.name).resolve()
        (self.root / "input").mkdir()
        (self.root / "output").mkdir()
        os.chdir(self.root)

    def tearDown(self):
        os.chdir(self.previous)
        self.temporary.cleanup()

    def csv(self, text, encoding="utf-8"):
        source = self.root / "input" / "原表.csv"
        source.write_bytes(text.encode(encoding))
        return source

    def run_cli(self, source, *arguments):
        return subprocess.run([
            sys.executable, str(SCRIPT), "--input", str(source),
            "--output", "output/result.md", *arguments,
        ], capture_output=True, text=True, encoding="utf-8",
            env={**os.environ, "PYTHONIOENCODING": "utf-8", "PYTHONDONTWRITEBYTECODE": "1"})

    def test_csv_preserves_literal_fields_and_original_record_numbers(self):
        source = self.csv('编号,备注,空值\r\n001,"甲|乙\n第二行",NA\r\n002,=1+1,\r\n')
        content, details = converter.csv_markdown(source, options())
        self.assertIn("| 原行号 | A | B | C |", content)
        self.assertIn("| 1 | 编号 | 备注 | 空值 |", content)
        self.assertIn("| 2 | 001 | 甲\\|乙<br>第二行 | NA |", content)
        self.assertIn("| 3 | 002 | =1+1 |  |", content)
        self.assertEqual(details["sheets"][0]["exportedRange"], "A1:C3")

    def test_csv_default_limit_and_explicit_later_window(self):
        source = self.csv("\n".join(",".join(f"r{row}c{column}" for column in range(1, 53))
                                    for row in range(1, 206)))
        content, details = converter.csv_markdown(source, options())
        self.assertEqual(details["sheets"][0]["exportedRange"], "A1:AX200")
        self.assertEqual(details["sheets"][0]["omittedRows"], 5)
        self.assertEqual(details["sheets"][0]["omittedColumns"], 2)
        self.assertIn("行：201–205；列：51–52", content)
        self.assertIn("不能据此声称已读取全表", content)
        self.assertNotIn("r201c1", content)
        later, details = converter.csv_markdown(source, options(start_row=201, start_column=51))
        self.assertEqual(details["sheets"][0]["exportedRange"], "AY201:AZ205")
        self.assertIn("| 205 | r205c51 | r205c52 |", later)

    def test_empty_or_outside_csv_window_does_not_invent_cells(self):
        source = self.csv("")
        content, details = converter.csv_markdown(source, options())
        self.assertEqual(details["sheets"][0]["rows"], 0)
        self.assertIsNone(details["sheets"][0]["exportedRange"])
        source = self.csv("a,b\nc,d")
        content, details = converter.csv_markdown(source, options(start_row=100, start_column=100))
        self.assertIsNone(details["sheets"][0]["exportedRange"])
        self.assertIn("行：1–2；列：1–2", content)
        self.assertNotIn("| 原行号", content)

    def test_csv_bom_and_explicit_legacy_encoding(self):
        source = self.csv("编号;名称\n001;原件", "utf-16")
        content, details = converter.csv_markdown(source, options())
        self.assertEqual(details["encoding"], "utf-16")
        self.assertEqual(details["delimiter"], ";")
        self.assertIn("| 2 | 001 | 原件 |", content)
        source = self.csv("编号,名称\n001,原件", "gb18030")
        content, details = converter.csv_markdown(source, options(encoding="gb18030"))
        self.assertIn("| 2 | 001 | 原件 |", content)

    @unittest.skipUnless(HAS_CHARDET, "chardet is supplied by the sandbox image")
    def test_csv_detects_legacy_chinese_encoding(self):
        source = self.csv("编号,名称\n" + "001,中文原始表格测试数据\n" * 30, "gb18030")
        content, _ = converter.csv_markdown(source, options())
        self.assertIn("中文原始表格测试数据", content)

    def test_cli_writes_new_result_preserves_original_and_rejects_overwrite(self):
        source = self.csv("编号,金额\n001,12.50")
        before = hashlib.sha256(source.read_bytes()).hexdigest()
        first = self.run_cli(source)
        self.assertEqual(first.returncode, 0, first.stderr)
        self.assertEqual(json.loads(first.stdout)["sheets"][0]["exportedRange"], "A1:B2")
        result = (self.root / "output/result.md").read_bytes()
        second = self.run_cli(source)
        self.assertNotEqual(second.returncode, 0)
        self.assertIn("Output already exists", second.stderr)
        self.assertEqual((self.root / "output/result.md").read_bytes(), result)
        self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), before)

    @unittest.skipUnless(HAS_OPENPYXL, "openpyxl is supplied by the sandbox image")
    def test_xlsx_sheets_formulas_merges_and_cached_zero(self):
        from openpyxl import Workbook

        source = self.root / "input" / "原表.xlsx"
        workbook = Workbook()
        first = workbook.active
        first.title = "原数据"
        first.append(["001", 0, False])
        first["A2"] = "=SUM(B1:B1)"
        first["B2"] = "=1+1"
        first["A3"] = "合并内容"
        first.merge_cells("A3:C3")
        second = workbook.create_sheet("隐藏明细")
        second.sheet_state = "hidden"
        second.append(["NA", "真实内容"])
        workbook.save(source)
        workbook.close()
        # A real OOXML cached zero distinguishes zero from an absent cached value.
        with zipfile.ZipFile(source) as archive:
            entries = {name: archive.read(name) for name in archive.namelist()}
        entries["xl/worksheets/sheet1.xml"] = entries["xl/worksheets/sheet1.xml"].replace(
            b"<f>SUM(B1:B1)</f><v></v>", b"<f>SUM(B1:B1)</f><v>0</v>")
        with zipfile.ZipFile(source, "w", zipfile.ZIP_DEFLATED) as archive:
            for name, value in entries.items():
                archive.writestr(name, value)
        before = source.read_bytes()
        content, details = converter.excel_markdown(source, options())
        self.assertFalse(details["libreOfficeFallback"])
        self.assertEqual([sheet["name"] for sheet in details["sheets"]], ["原数据", "隐藏明细"])
        self.assertIn("| 1 | 001 | 0 | FALSE |", content)
        self.assertIn("=SUM(B1:B1)<br>[缓存值：0]", content)
        self.assertIn("=1+1<br>[无缓存值；未计算]", content)
        self.assertIn("A3:C3", content)
        self.assertIn("| 3 | 合并内容 |  |  |", content)
        self.assertIn("工作表状态：hidden", content)
        self.assertEqual(source.read_bytes(), before)

    @unittest.skipUnless(HAS_OPENPYXL, "openpyxl is supplied by the sandbox image")
    def test_xlsx_window_keeps_coordinates_and_does_not_fill_merge_anchor_outside(self):
        from openpyxl import Workbook

        source = self.root / "input" / "大表.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet["A1"] = "ANCHOR-OUTSIDE-SELECTED-RANGE"
        sheet.merge_cells("A1:C2")
        sheet["AZ205"] = "最后一格"
        workbook.save(source)
        workbook.close()
        content, details = converter.excel_markdown(source, options(start_row=2, start_column=2))
        self.assertEqual(details["sheets"][0]["exportedRange"], "B2:AY201")
        self.assertIn("锚格在范围外", content)
        self.assertNotIn("ANCHOR-OUTSIDE-SELECTED-RANGE", content)
        self.assertNotIn("最后一格", content)
        later, details = converter.excel_markdown(source, options(start_row=205, start_column=52))
        self.assertEqual(details["sheets"][0]["exportedRange"], "AZ205:AZ205")
        self.assertIn("| 205 | 最后一格 |", later)

    @unittest.skipUnless(HAS_OPENPYXL and LIBREOFFICE, "LibreOffice is supplied by the sandbox image")
    def test_libreoffice_fallback_uses_scratch_copy_and_discloses_recalculation(self):
        from openpyxl import Workbook

        source = self.root / "input" / "原表.xlsx"
        workbook = Workbook()
        workbook.active.append(["原件", "=1+1"])
        workbook.save(source)
        workbook.close()
        before = source.read_bytes()
        original = converter.xlsx_sheets

        def primary_fails_only_for_original(path, settings):
            if path == source:
                raise ValueError("synthetic primary reader failure")
            return original(path, settings)

        with patch.object(converter, "xlsx_sheets", side_effect=primary_fails_only_for_original):
            content, details = converter.excel_markdown(source, options())
        self.assertTrue(details["libreOfficeFallback"])
        self.assertIn("以下缓存值来自转换副本", content)
        self.assertIn("原件", content)
        self.assertEqual(source.read_bytes(), before)
        self.assertEqual(list((self.root / "scratch").iterdir()), [])

    @unittest.skipUnless(HAS_OPENPYXL and HAS_PANDAS and HAS_XLRD and LIBREOFFICE,
                         "pandas/xlrd/LibreOffice are supplied by the sandbox image")
    def test_real_xls_reads_multiple_sheets_without_losing_text_or_error_cells(self):
        from openpyxl import Workbook

        original = self.root / "input" / "legacy.xlsx"
        workbook = Workbook()
        first = workbook.active
        first.append(["001", "NA", "#DIV/0!"])
        first["A2"] = "合并"
        first.merge_cells("A2:B2")
        workbook.create_sheet("Second").append([0, False])
        workbook.save(original)
        workbook.close()
        subprocess.run([LIBREOFFICE, "-env:UserInstallation=" + (self.root / "lo-profile").as_uri(),
                        "--headless", "--convert-to", "xls:MS Excel 97", "--outdir",
                        str(self.root / "input"), str(original)], check=True, capture_output=True, timeout=60)
        source = original.with_suffix(".xls")
        before = source.read_bytes()
        content, details = converter.excel_markdown(source, options())
        self.assertFalse(details["libreOfficeFallback"])
        self.assertEqual(len(details["sheets"]), 2)
        self.assertIn("| 1 | 001 | NA | #DIV/0! |", content)
        self.assertIn("A2:B2", content)
        self.assertIn("不能还原公式表达式", content)
        self.assertEqual(source.read_bytes(), before)


if __name__ == "__main__":
    unittest.main(verbosity=2)
