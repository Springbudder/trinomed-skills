#!/usr/bin/env python3
"""Offline structural document extraction; no inference or business operations."""

import argparse
import codecs
import csv
from datetime import date, datetime, time
import html
import json
from pathlib import Path
import subprocess
import tempfile


def new_output(value):
    target = Path(value).resolve()
    target.relative_to((Path.cwd() / "output").resolve())
    target.parent.mkdir(parents=True, exist_ok=True)
    if target.exists():
        raise FileExistsError("Output already exists; choose a new version name.")
    return target


def pdf_markdown(source):
    info = subprocess.run(["pdfinfo", str(source)], check=True, capture_output=True, text=True)
    page_values = [line.partition(":")[2].strip() for line in info.stdout.splitlines()
                   if line.partition(":")[0].strip() == "Pages"]
    if len(page_values) != 1:
        raise ValueError("PDF page count was not reported by pdfinfo.")
    page_count = int(page_values[0])
    sections = ["PDF 支持说明：以下为逐页文字抽取，固定宽度块保留抽取到的行列空白；"
                "不保证多栏阅读顺序、表格单元格、图片或原版式还原。"]
    empty_pages = []
    for number in range(1, page_count + 1):
        result = subprocess.run([
            "pdftotext", "-f", str(number), "-l", str(number), "-layout", "-enc", "UTF-8",
            str(source), "-",
        ], check=True, capture_output=True, text=True, encoding="utf-8")
        text = result.stdout.replace("\f", "").strip("\r\n")
        sections.append(f"<!-- source-page: {number} -->\n## 第 {number} 页")
        if text.strip():
            sections.append("\n".join("    " + line for line in text.splitlines()))
        else:
            empty_pages.append(number)
            sections.append("[本页没有抽取到文字；可能为扫描页或空白页，请结合原页检查，必要时逐页 OCR。]")
    return "\n\n".join(sections), {"pages": page_count, "pagesWithoutExtractedText": empty_pages}


def cell_text(value):
    return value.replace("\\", "\\\\").replace("|", "\\|").replace("\n", " / ")


def docx_markdown(source):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(str(source))
    sections = ["DOCX 支持说明：保留正文段落与表格顺序；表格首行用作 Markdown 表头。"
                "不保证浮动文本框、页眉页脚、批注、复杂编号、嵌套表与图片还原；DOCX 不估算页数。"]
    if document.inline_shapes:
        sections.append(f"[原文含 {len(document.inline_shapes)} 个内联图形，本次文字转换未包含这些图像。]")
    table_count = 0
    for element in document.element.body.iterchildren():
        if element.tag == qn("w:p"):
            paragraph = Paragraph(element, document)
            text = paragraph.text
            if not text.strip():
                continue
            style_name = paragraph.style.name if paragraph.style is not None else ""
            heading_level = style_name.removeprefix("Heading ") if style_name.startswith("Heading ") else ""
            if heading_level in {"1", "2", "3", "4", "5", "6"}:
                sections.append("#" * int(heading_level) + " " + text)
            else:
                sections.append(text)
        elif element.tag == qn("w:tbl"):
            table = Table(element, document)
            rows = [[cell_text(cell.text) for cell in row.cells] for row in table.rows]
            if not rows or not rows[0]:
                continue
            table_count += 1
            merged = any(child.tag in {qn("w:gridSpan"), qn("w:vMerge")}
                         for child in element.iter())
            if merged:
                sections.append("[此表含合并单元格，以下展开为单元格矩阵；源合并区域的文字可能重复。]")
            columns = max(len(row) for row in rows)
            rows = [row + [""] * (columns - len(row)) for row in rows]
            lines = ["| " + " | ".join(rows[0]) + " |", "| " + " | ".join(["---"] * columns) + " |"]
            lines.extend("| " + " | ".join(row) + " |" for row in rows[1:])
            sections.append("\n".join(lines))
    return "\n\n".join(sections), {"tables": table_count, "pages": None}


def column_name(number):
    result = ""
    while number:
        number, digit = divmod(number - 1, 26)
        result = chr(65 + digit) + result
    return result


def spreadsheet_text(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        value = value.isoformat()
    elif isinstance(value, bool):
        value = "TRUE" if value else "FALSE"
    # Cell content is literal text; Markdown punctuation must not change its meaning.
    value = html.escape(str(value), quote=False).replace("\\", "\\\\")
    for character in "|`*_[]":
        value = value.replace(character, "\\" + character)
    return value.replace("\r\n", "\n").replace("\r", "\n").replace("\n", "<br>")


def table_window(total_rows, total_columns, options):
    rows = range(options.start_row, min(total_rows + 1, options.start_row + options.max_rows))
    columns = range(options.start_column, min(total_columns + 1, options.start_column + options.max_columns))
    exported = (f"{column_name(columns.start)}{rows.start}:"
                f"{column_name(columns.stop - 1)}{rows.stop - 1}") if rows and columns else None
    return rows, columns, {
        "rows": total_rows, "columns": total_columns, "exportedRange": exported,
        "omittedRows": total_rows - (len(rows) if columns else 0),
        "omittedColumns": total_columns - (len(columns) if rows else 0),
    }


def render_sheet(name, total_rows, total_columns, values, options, notes=()):
    rows, columns, details = table_window(total_rows, total_columns, options)
    details["name"] = name
    sections = [f"## 工作表：{spreadsheet_text(name)}",
                f"记录范围：{total_rows} 行 × {total_columns} 列；"
                f"实际导出范围：{details['exportedRange'] or '无（空表或请求范围之外）'}。"]
    if details["omittedRows"] or details["omittedColumns"]:
        omitted = []
        for label, total, selected in (("行", total_rows, rows), ("列", total_columns, columns)):
            effective = selected if rows and columns else range(1, 1)
            spans = []
            if not effective and total:
                spans.append(f"1–{total}")
            elif effective:
                if effective.start > 1:
                    spans.append(f"1–{effective.start - 1}")
                if effective.stop <= total:
                    spans.append(f"{effective.stop}–{total}")
            if spans:
                omitted.append(f"{label}：" + "、".join(spans))
        sections.append("[本段仅含所列范围，已截断/省略 " + "；".join(omitted)
                        + "。不能据此声称已读取全表或完成全表统计；可指定后续区间继续读取。]")
    sections.extend(notes)
    if rows and columns:
        lines = ["| 原行号 | " + " | ".join(column_name(number) for number in columns) + " |",
                 "| --- | " + " | ".join("---" for _ in columns) + " |"]
        lines.extend("| " + str(number) + " | " + " | ".join(row) + " |"
                     for number, row in zip(rows, values))
        sections.append("\n".join(lines))
    return "\n\n".join(sections), details


def xlsx_sheets(source, options):
    from openpyxl import load_workbook

    workbook = load_workbook(source, data_only=False)
    cached = load_workbook(source, data_only=True, read_only=True)
    sections, sheets = [], []
    try:
        for sheet in workbook.worksheets:
            rows, columns, _ = table_window(sheet.max_row, sheet.max_column, options)
            cached_rows = cached[sheet.title].iter_rows(
                min_row=rows.start, max_row=rows.stop - 1,
                min_col=columns.start, max_col=columns.stop - 1,
            ) if rows and columns else ()
            values, formulas = [], 0
            for row_number, cached_row in zip(rows, cached_rows):
                row = []
                for column, cached_cell in zip(columns, cached_row):
                    cell = sheet.cell(row_number, column)
                    value = spreadsheet_text(cell.value)
                    if cell.data_type == "f":
                        formulas += 1
                        expression = getattr(cell.value, "text", cell.value)
                        value = spreadsheet_text(expression)
                        value += ("<br>[缓存值：" + spreadsheet_text(cached_cell.value) + "]"
                                  if cached_cell.value is not None else "<br>[无缓存值；未计算]")
                    row.append(value)
                values.append(row)
            merges = [str(area) for area in sheet.merged_cells.ranges
                      if rows and columns and area.min_row < rows.stop and area.max_row >= rows.start
                      and area.min_col < columns.stop and area.max_col >= columns.start]
            notes = [f"工作表状态：{sheet.sheet_state}（隐藏行列也按原坐标读取）。",
                     f"本段公式单元格：{formulas} 个；保留公式表达式，缓存值可能过期，不计算公式。"]
            if merges:
                notes.append("[本段相交合并区域：" + "、".join(sorted(merges))
                             + "。只保留原锚格值，不重复填充；锚格在范围外时本段不含其值。]")
            section, details = render_sheet(sheet.title, sheet.max_row, sheet.max_column,
                                            values, options, notes)
            details.update({"formulaCellsInRange": formulas, "mergedRangesInRange": sorted(merges)})
            sections.append(section)
            sheets.append(details)
    finally:
        workbook.close()
        cached.close()
    return sections, sheets


def xls_sheets(source, options):
    import pandas as pd
    import xlrd

    sections, sheets = [], []
    with pd.ExcelFile(source, engine="xlrd", engine_kwargs={"formatting_info": True}) as workbook:
        for name in workbook.sheet_names:
            sheet = workbook.book.sheet_by_name(name)
            rows, columns, _ = table_window(sheet.nrows, sheet.ncols, options)
            values = []
            if rows and columns:
                frame = workbook.parse(name, header=None, dtype=object, na_filter=False,
                                       skiprows=rows.start - 1, nrows=len(rows))
                for offset, row_number in enumerate(rows):
                    values_row = []
                    for column in columns:
                        cell = sheet.cell(row_number - 1, column - 1)
                        if cell.ctype == xlrd.XL_CELL_ERROR:
                            value = xlrd.biffh.error_text_from_code[cell.value]
                        else:
                            value = (frame.iat[offset, column - 1]
                                     if offset < len(frame.index) and column <= len(frame.columns) else None)
                        values_row.append(spreadsheet_text(value))
                    values.append(values_row)
            merges = [f"{column_name(left + 1)}{top + 1}:{column_name(right)}{bottom}"
                      for top, bottom, left, right in sheet.merged_cells
                      if rows and columns and top + 1 < rows.stop and bottom >= rows.start
                      and left + 1 < columns.stop and right >= columns.start]
            notes = ["XLS 读取保存时的单元格值；此读取器不能还原公式表达式或确认公式缓存是否最新。",
                     f"工作表隐藏状态代码：{sheet.visibility}（0 为可见；隐藏行列也按原坐标读取）。"]
            if merges:
                notes.append("[本段相交合并区域：" + "、".join(merges)
                             + "。只保留原锚格值，不重复填充。]")
            section, details = render_sheet(name, sheet.nrows, sheet.ncols, values, options, notes)
            details["mergedRangesInRange"] = merges
            sections.append(section)
            sheets.append(details)
    return sections, sheets


def excel_markdown(source, options):
    reader = xlsx_sheets if source.suffix.lower() == ".xlsx" else xls_sheets
    fallback = False
    try:
        sections, sheets = reader(source, options)
    except Exception as primary_error:
        # LibreOffice only writes an isolated scratch copy; the original remains read-only.
        scratch = Path.cwd() / "scratch"
        scratch.mkdir(exist_ok=True)
        with tempfile.TemporaryDirectory(prefix="document-markdown-", dir=scratch) as directory:
            temporary = Path(directory)
            subprocess.run([
                "libreoffice", "-env:UserInstallation=" + (temporary / "profile").as_uri(),
                "--headless", "--convert-to", "xlsx", "--outdir", str(temporary), str(source),
            ], check=True, capture_output=True, text=True, timeout=60)
            converted = temporary / (source.stem + ".xlsx")
            if not converted.is_file():
                raise ValueError("Excel reader failed and LibreOffice produced no XLSX copy.") from primary_error
            sections, sheets = xlsx_sheets(converted, options)
            fallback = True
    introduction = ("Excel 支持说明：每个工作表独立成段；列字母与原行号是坐标，原首行仍为数据，"
                    "不猜测业务表头。范围来自工作表记录尺寸，可能包含仅有格式的空行列。"
                    "读取单元格值，不复刻数字显示格式、图表、图片、批注或原版式；单元格换行以 <br> 表示。")
    if fallback:
        introduction += ("\n\n[主读取器未能读取，已用 LibreOffice 在 scratch 中转换副本。"
                         "转换可能重新计算公式或改变格式；以下缓存值来自转换副本，不代表原件保存时的值。原件未修改。]")
    return "\n\n".join([introduction, *sections]), {"sheets": sheets, "libreOfficeFallback": fallback, "pages": None}


def csv_markdown(source, options):
    encoding = options.encoding
    if not encoding:
        with source.open("rb") as stream:
            sample = stream.read(65536)
        if sample.startswith(codecs.BOM_UTF8):
            encoding = "utf-8-sig"
        elif sample.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
            encoding = "utf-16"
        else:
            try:
                # A partial final UTF-8 character in the sample is not an encoding failure.
                codecs.getincrementaldecoder("utf-8")().decode(sample, final=False)
                encoding = "utf-8"
            except UnicodeDecodeError:
                import chardet
                encoding = chardet.detect(sample)["encoding"]
                if not encoding:
                    raise ValueError("CSV encoding could not be detected; specify --encoding.")
    with source.open("r", encoding=encoding, newline="") as stream:
        sample = stream.read(65536)
        stream.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
            delimiter = dialect.delimiter
        except csv.Error:
            delimiter = ","
        values, total_rows, total_columns = [], 0, 0
        for number, row in enumerate(csv.reader(stream, delimiter=delimiter), 1):
            total_rows = number
            total_columns = max(total_columns, len(row))
            if options.start_row <= number < options.start_row + options.max_rows:
                values.append(row[options.start_column - 1:options.start_column - 1 + options.max_columns])
    _, columns, _ = table_window(total_rows, total_columns, options)
    rendered = [[spreadsheet_text(value) for value in row] + [""] * (len(columns) - len(row)) for row in values]
    section, details = render_sheet(source.name, total_rows, total_columns, rendered, options)
    introduction = (f"CSV 支持说明：编码 {encoding}；分隔符 {repr(delimiter)}；作为一个工作表读取。"
                    "字段按原文保留，不推断表头、数字或公式；列字母与原行号是记录坐标（引号内换行不增加记录号）。"
                    "编码与分隔符自动检测可能不准确，须核对原件；必要时用 --encoding 指定编码。"
                    "单元格换行以 <br> 表示。")
    return introduction + "\n\n" + section, {"sheets": [details], "encoding": encoding, "delimiter": delimiter, "pages": None}


def positive_integer(value):
    number = int(value)
    if number < 1:
        raise argparse.ArgumentTypeError("Value must be a positive integer.")
    return number


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--start-row", type=positive_integer, default=1)
    parser.add_argument("--start-column", type=positive_integer, default=1)
    parser.add_argument("--max-rows", type=positive_integer, default=200)
    parser.add_argument("--max-columns", type=positive_integer, default=50)
    parser.add_argument("--encoding", help="CSV text encoding; otherwise detect BOM/UTF-8/chardet.")
    arguments = parser.parse_args()
    source = Path(arguments.input).resolve(strict=True)
    target = new_output(arguments.output)
    extension = source.suffix.lower()
    if extension == ".pdf":
        content, details = pdf_markdown(source)
    elif extension == ".docx":
        content, details = docx_markdown(source)
    elif extension in {".xlsx", ".xls"}:
        content, details = excel_markdown(source, arguments)
    elif extension == ".csv":
        content, details = csv_markdown(source, arguments)
    elif extension in {".txt", ".md"}:
        content, details = source.read_text(encoding="utf-8"), {"pages": None}
    else:
        raise ValueError("Supported files: PDF, DOCX, XLSX, XLS, CSV, UTF-8 TXT and Markdown.")
    with target.open("x", encoding="utf-8", newline="\n") as stream:
        stream.write(content + "\n")
    print(json.dumps({"output": str(target.relative_to(Path.cwd())),
                      "byteSize": target.stat().st_size, **details}, ensure_ascii=False))


if __name__ == "__main__":
    main()
